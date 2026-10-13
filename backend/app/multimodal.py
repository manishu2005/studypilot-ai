from dataclasses import dataclass
from pathlib import Path
import pymupdf
from io import BytesIO
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

@dataclass
class ExtractedImage:
    source_type:str
    location_number:int
    image_index:int
    extension:str
    data:bytes
    width:int
    height:int

def extract_pdf_images(data:bytes)->list[ExtractedImage]:
    images = []
    with pymupdf.open(stream=data, filetype="pdf")as document:
        for page_index, page in enumerate(document):
            page_number = page_index + 1
            image_list = page.get_images(full=True)
            seen_xrefs=set()
            for image_index, image_info in enumerate(image_list, start=1):
                xref = image_info[0]
                if xref in seen_xrefs:
                    continue

                seen_xrefs.add(xref)
                extracted = document.extract_image(xref)
                if not extracted:
                    continue
                image_data = extracted["image"]
                extension = extracted["ext"]
                images.append(
                    ExtractedImage(
                        source_type="pdf",
                        location_number=page_number,
                        image_index=image_index,
                        extension=extension,
                        data=image_data,
                        width=extracted["width"],
                        height=extracted["height"],
                    )
                )

    return images

def extract_pptx_images(data:bytes)->list[ExtractedImage]:
    images = []
    presentation = Presentation(BytesIO(data))
    for slide_number, slide in enumerate(presentation.slides, start=1):
        image_index=0
        for shape in slide.shapes:
            if(shape.shape_type != MSO_SHAPE_TYPE.PICTURE):
                continue

            image_index += 1
            image = shape.image
            images.append(
                ExtractedImage(
                    source_type="pptx",
                    location_number=slide_number,
                    image_index=image_index,
                    extension=image.ext,
                    data=image.blob,
                    width=image.size[0],
                    height=image.size[1],
                )
            )
    return images

def extract_images(filename:str, data:bytes)->list[ExtractedImage]:
    extension = Path(filename).suffix.lower()
    if extension == ".pdf":
        return extract_pdf_images(data)

    if extension == "pptx":
        return extract_pptx_images(data)

    raise ValueError(f"Image extraction is not supported for {extension} files")

