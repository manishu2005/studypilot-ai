from fastapi.testclient import TestClient
from app.main import app, get_service
import pymupdf
from app.ingestion import extract_document, extract_pdf, extract_pptx, split_text
from pptx import Presentation
import pytest

class FakeStudyService:
    def __init__(self):
        self.documents = []

    def list_documents(self):
        return self.documents

    def ingest_pdf(self, filename, data):
        document = {
            "id": "test-document-001",
            "document_id": "test-document-001",
            "filename": filename,
            "status": "ready",
            "chunk_count": 1,
            "pages_needing_ocr": [],
        }
        self.documents.append(document)
        return document

    def search(self, query, top_k=5):
        return [
            {
                "text": "BFS uses a queue.",
                "source": "sample.pdf",
                "document_id": "test-document-001",
                "page_number": 2,
                "chunk_index": 0,
                "distance": 0.15,
            }
        ]



@pytest.fixture
def client():
    fake_service = FakeStudyService()
    def override_service():
        return fake_service

    app.dependency_overrides[get_service] = override_service
    with TestClient(app) as test_client:
        yield test_client

    app.dependency_overrides.pop(get_service, None)

def test_health(client):
    response = client.get("/health")

    assert response.status_code == 200
    assert response.json()["status"] == "healthy"


def test_upload_rejects_non_pdf(client):
    response = client.post(
        "/api/documents/upload",
        files={
            "file": (
                "notes.txt",
                b"not a pdf",
                "text/plain",
            )
        },
    )

    assert response.status_code == 415


def test_upload_rejects_empty_file(client):
    response = client.post(
        "/api/documents/upload",
        files={
            "file": (
                "notes.pdf",
                b"",
                "application/pdf",
            )
        },
    )

    assert response.status_code == 400


def test_search_returns_source_location(client):
    response = client.post(
        "/api/search",
        json={
            "query": "What data structure does BFS use?",
            "top_k": 3,
        },
    )

    assert response.status_code == 200

    result = response.json()["matches"][0]

    assert result["text"] == "BFS uses a queue."
    assert result["source"] == "sample.pdf"
    assert result["page_number"] == 2


def test_empty_search_query_rejected(client):
    response = client.post(
        "/api/search",
        json={"query": "", "top_k": 5},
    )

    assert response.status_code == 422

def test_pdf_extraction_preserves_page_number():
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text(
        (72,72), "BFS uses a queue to traverse a graph"
    )

    pdf_bytes = pdf.tobytes()
    pdf.close()

    result = extract_pdf(pdf_bytes)

    assert result.chunks
    assert result.chunks[0].page_number == 1
    assert "queue" in result.chunks[0].text
    assert result.pages_needing_ocr == []

def test_split_text_creates_overlapping_chunks():
    chunks=split_text(
        "A B C D E F G H I J K L M N O P Q R S T U V W X Y Z",
        chunk_size=10,
        overlap=3,
    )

    assert len(chunks) == 4
    assert chunks[0] == "A B C D E F G H I J"
    assert chunks[1] == "H I J K L M N O P Q"
    assert chunks[2] == "O P Q R S T U V W X"
    assert chunks[3] == "V W X Y Z"

def test_pdf_extraction_preserves_page_count():

    pdf = pymupdf.open()

    page1 = pdf.new_page()

    page1.insert_text(
        (72, 72),
        "BFS uses a queue.",
    )

    page2 = pdf.new_page()

    page2.insert_text(
        (72, 72),
        "DFS uses a stack.",
    )

    pdf_bytes = pdf.tobytes()

    pdf.close()

    result = extract_pdf(pdf_bytes)

    assert result.page_count == 2

    assert len(result.chunks) == 2

    assert result.chunks[0].page_number == 1

    assert result.chunks[1].page_number == 2

    assert result.pages_needing_ocr == []

def test_scanned_like_pdf_is_flagged_for_ocr():
    pdf = pymupdf.open()

    page = pdf.new_page()

    page.draw_rect(pymupdf.Rect(50,50,300,150))
    pdf_bytes = pdf.tobytes()
    pdf.close()
    result = extract_pdf(pdf_bytes)
    assert result.page_count == 1
    assert result.pages_needing_ocr == [1]

def test_pptx_extraction_preserves_slide_numbers(tmp_path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text="Breadth First Search"
    slide.placeholders[1].text = "BFS uses a queue"
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text="Depth First Search"
    slide.placeholders[1].text = "DFS uses a stack"
    pptx_path = tmp_path / "test.pptx"
    presentation.save(pptx_path)
    data = pptx_path.read_bytes()
    result = extract_pptx(data)
    assert result.location_count==2
    assert len(result.chunks)==2
    assert result.chunks[0].location_number == 1
    assert result.chunks[1].location_number == 2
    assert (result.chunks[0].source_type == "pptx")
    assert ("queue" in result.chunks[0].text.lower())
    assert ("stack" in result.chunks[1].text.lower())

def test_unsupported_document_type_rejected():

    with pytest.raises(
        ValueError,
        match="Unsupported document type",
    ):

        extract_document(
            "notes.txt",
            b"hello world",
        )