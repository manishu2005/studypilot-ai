from dataclasses import dataclass
from io import BytesIO

import pymupdf
from pptx import Presentation
from pydantic import BaseModel

class SourceChunk(BaseModel):
    text: str
    page_number: int
    chunk_index: int
    source_type: str


@dataclass
class PDFExtractionResult:
    chunks: list[SourceChunk]
    pages_needing_ocr: list[int]
    page_count: int

    def __iter__(self):
        # Allows:
        # chunks, pages_needing_ocr = extract_pdf(data)
        yield self.chunks
        yield self.pages_needing_ocr


# ============================================================
# PPTX
# ============================================================

class PPTXChunk(BaseModel):
    text: str
    location_number: int
    chunk_index: int
    source_type: str


@dataclass
class PPTXExtractionResult:
    chunks: list[PPTXChunk]
    location_count: int


# ============================================================
# TEXT SPLITTING
# ============================================================

def split_text(
    text: str,
    chunk_size: int = 900,
    overlap: int = 120
) -> list[str]:

    text = " ".join(text.split())

    if not text:
        return []

    if chunk_size <= 0:
        raise ValueError(
            "chunk_size must be a positive integer."
        )

    if overlap < 0 or overlap >= chunk_size:
        raise ValueError(
            "overlap must be non-negative and smaller than chunk_size"
        )

    words = text.split()

    chunks = []
    start = 0

    while start < len(words):

        end = min(
            start + chunk_size,
            len(words)
        )

        chunk = " ".join(words[start:end]).strip()

        if chunk:
            chunks.append(chunk)

        if end >= len(words):
            break

        start = end - overlap

    return chunks

# ============================================================
# PDF EXTRACTION
# ============================================================

def extract_pdf(data: bytes) -> PDFExtractionResult:

    if not data.startswith(b"%PDF"):
        raise ValueError(
            "The uploaded file is not a valid PDF."
        )

    chunks: list[SourceChunk] = []
    pages_needing_ocr: list[int] = []

    try:

        with pymupdf.open(
            stream=data,
            filetype="pdf"
        ) as pdf:

            page_count = len(pdf)

            if page_count == 0:
                raise ValueError(
                    "The PDF contains no pages."
                )

            for page_index, page in enumerate(pdf):

                page_number = page_index + 1

                text = page.get_text(
                    "text",
                    sort=True
                ).strip()

                if not text:
                    pages_needing_ocr.append(
                        page_number
                    )
                    continue

                page_chunks = split_text(text)

                for chunk_index, chunk_text in enumerate(
                    page_chunks
                ):

                    chunks.append(
                        SourceChunk(
                            text=chunk_text,
                            page_number=page_number,
                            chunk_index=chunk_index,
                            source_type="pdf"
                        )
                    )

    except ValueError:
        raise

    except Exception as exc:
        raise ValueError(
            "The PDF could not be opened or processed."
        ) from exc

    if not chunks and not pages_needing_ocr:
        raise ValueError(
            "No extractable text was found in the PDF."
        )

    return PDFExtractionResult(
        chunks=chunks,
        pages_needing_ocr=pages_needing_ocr,
        page_count=page_count
    )


# ============================================================
# PPTX EXTRACTION
# ============================================================

def extract_pptx(data: bytes) -> PPTXExtractionResult:

    if not data:
        raise ValueError(
            "The PowerPoint file is empty or invalid."
        )

    try:

        presentation = Presentation(
            BytesIO(data)
        )

    except Exception as exc:
        raise ValueError(
            "The PowerPoint file could not be opened or processed."
        ) from exc

    slide_count = len(
        presentation.slides
    )

    if slide_count == 0:
        raise ValueError(
            "The PowerPoint presentation contains no slides."
        )

    chunks: list[PPTXChunk] = []

    for slide_index, slide in enumerate(
        presentation.slides
    ):

        slide_number = slide_index + 1

        text_parts = []

        for shape in slide.shapes:

            if not shape.has_text_frame:
                continue

            text = shape.text.strip()

            if text:
                text_parts.append(text)

        slide_text = "\n".join(
            text_parts
        ).strip()

        if not slide_text:
            continue

        slide_chunks = split_text(
            slide_text
        )

        for chunk_index, chunk_text in enumerate(
            slide_chunks
        ):

            chunks.append(
                PPTXChunk(
                    text=chunk_text,
                    location_number=slide_number,
                    chunk_index=chunk_index,
                    source_type="pptx"
                )
            )

    return PPTXExtractionResult(
        chunks=chunks,
        location_count=slide_count
    )


# ============================================================
# DOCUMENT DISPATCHER
# ============================================================

def extract_document(
    filename: str,
    data: bytes
):

    extension = filename.lower().rsplit(
        ".",
        1
    )[-1]

    if extension == "pdf":
        return extract_pdf(data)

    if extension == "pptx":
        return extract_pptx(data)

    raise ValueError(
        "Unsupported document type. "
        "Supported formats: PDF and PPTX."
    )