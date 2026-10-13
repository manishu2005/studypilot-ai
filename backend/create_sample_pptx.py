from pptx import Presentation
presentation = Presentation()

slide = presentation.slides.add_slide(presentation.slide_layouts[1])
slide.shapes.title.text= "BFS"
slide.placeholders[1].text = (
    "BFS traverses a graph level by level.\n"
    "It commonly uses a queue.\n"
    "The queue follows FIFO order."
)

slide = presentation.slides.add_slide(presentation.slide_layouts[1])
slide.shapes.title.text="DFS"
slide.placeholders[1].text = (
    "DFS explores a graph deeply before "
    "backtracking.\n"
    "DFS can use a stack."
)

slide = presentation.slides.add_slide(presentation.slide_layouts[1])
slide.shapes.title.text="BFS vs DFS"
slide.placeholders[1].text = (
    "BFS uses a queue.\n"
    "DFS commonly uses a stack.\n"
    "Both are graph traversal algorithms."
)

presentation.save("sample_lecture.pptx")
print("Created")